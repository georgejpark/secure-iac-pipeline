#!/usr/bin/env python3
"""Yellow-highlight everything spoken in docs/word/1-RUNBOOK.docx.

The rule the runbook states: yellow = say it. Applied after each export:
  - Part B, all of it (it is the spoken introduction)
  - Part C and D: every quote box; any line containing a quoted phrase;
    all of STEP 9; the "What to say" column of the walkthrough table;
    the "three things to land"; the Answer column of the questions table
Code blocks are never highlighted.
"""
import re, sys
from docx import Document
from docx.enum.text import WD_COLOR_INDEX
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

path = sys.argv[1] if len(sys.argv) > 1 else "docs/word/1-RUNBOOK.docx"
d = Document(path)
CODE = ("Source Code", "Verbatim Char")
part = None; section = ""; n = 0

def mark(p):
    global n
    for r in p.runs: r.font.highlight_color = WD_COLOR_INDEX.YELLOW
    n += 1

for child in d.element.body.iterchildren():
    if child.tag == qn('w:p'):
        p = Paragraph(child, d); st = p.style.name; t = p.text
        if st.startswith("Heading") or st == "Title":
            m = re.search(r"PART\s+([A-E])\b", t)
            if m: part = m.group(1)
            section = t
            continue
        if st in CODE or not t.strip():
            continue
        spoken = (
            part == "B"
            or (part in ("C", "D") and st == "Block Text")
            or (part in ("C", "D") and '"' in t)
            or section.startswith("STEP 9")
            or section.startswith(("The three things to land", "The closing line"))
            or section.startswith("What that script does")
        )
        if spoken: mark(p)
    elif child.tag == qn('w:tbl'):
        tb = Table(child, d)
        hdr = [c.text.strip() for c in tb.rows[0].cells]
        if part == "B":
            for row in tb.rows:
                for c in row.cells:
                    for p in c.paragraphs: mark(p)
        elif "What to say" in hdr or "Answer" in hdr:
            col = hdr.index("What to say") if "What to say" in hdr else hdr.index("Answer")
            for row in tb.rows[1:]:
                for p in row.cells[col].paragraphs: mark(p)
d.save(path)
print(f"highlighted {n} blocks in {path}")
